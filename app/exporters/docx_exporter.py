from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt

from PIL import Image

from app.config import BASE_DIR
from app.exporters.common import sanitize_text


LOGO_PATH = (
    BASE_DIR
    / "assets"
    / "legalease_logo.png"
)


def _set_default_font(document: Document):

    normal_style = document.styles["Normal"]

    normal_style.font.name = "Times New Roman"
    normal_style.font.size = Pt(12)


def _get_valid_logo():

    """
    Validate the LegalEase logo using Pillow.

    Returns a BytesIO PNG stream if valid.
    Returns None if the logo is missing or invalid.
    """

    if not LOGO_PATH.exists():

        print(
            f"Logo not found: {LOGO_PATH}"
        )

        return None

    try:

        # Open and validate image.
        with Image.open(LOGO_PATH) as image:

            image.verify()

        # Re-open because verify() closes validation state.
        with Image.open(LOGO_PATH) as image:

            image = image.convert("RGBA")

            output = BytesIO()

            image.save(
                output,
                format="PNG"
            )

            output.seek(0)

            return output

    except Exception as error:

        print(
            "Invalid LegalEase logo. "
            f"Logo will be skipped. Error: {error}"
        )

        return None


def _add_logo(document: Document):

    """
    Add logo only if it is a valid image.

    Invalid/missing logo will NOT stop DOCX generation.
    """

    logo_stream = _get_valid_logo()

    if logo_stream is None:

        return

    try:

        paragraph = document.add_paragraph()

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run()

        run.add_picture(
            logo_stream,
            width=Inches(2.0)
        )

    except Exception as error:

        print(
            "Could not insert logo. "
            f"Continuing without logo: {error}"
        )


def _add_footer(document: Document):

    for section in document.sections:

        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            "LegalEase - Drafting aid only. "
            "Review with a qualified legal "
            "professional before use."
        )

        run.font.name = (
            "Times New Roman"
        )

        run.font.size = Pt(8)


def _add_title(
    document: Document,
    doc_type: str
):

    paragraph = document.add_paragraph()

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    run = paragraph.add_run(
        doc_type.upper()
    )

    run.bold = True

    run.font.name = (
        "Times New Roman"
    )

    run.font.size = Pt(16)


def _add_terms_table(
    document: Document,
    terms: str
):

    if not terms:

        return

    items = []

    for item in terms.split(";"):

        item = item.strip()

        if item:

            items.append(item)

    if not items:

        return

    document.add_paragraph()

    heading = document.add_paragraph()

    heading_run = heading.add_run(
        "Key Terms"
    )

    heading_run.bold = True

    heading_run.font.name = (
        "Times New Roman"
    )

    heading_run.font.size = Pt(12)

    table = document.add_table(
        rows=1,
        cols=2
    )

    table.style = "Table Grid"

    table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    header = table.rows[0].cells

    header[0].text = "No."

    header[1].text = "Term"

    for index, item in enumerate(
        items,
        start=1
    ):

        cells = table.add_row().cells

        cells[0].text = str(index)

        cells[1].text = item


def create_docx(
    text: str,
    doc_type: str,
    terms: str = "",
) -> bytes:

    """
    Create a professionally formatted DOCX file.

    The logo is optional. Invalid logos are ignored.
    """

    document = Document()

    _set_default_font(document)

    # ------------------------------------------------------
    # LOGO
    # ------------------------------------------------------

    _add_logo(document)

    # ------------------------------------------------------
    # TITLE
    # ------------------------------------------------------

    _add_title(
        document,
        doc_type
    )

    document.add_paragraph()

    # ------------------------------------------------------
    # DOCUMENT CONTENT
    # ------------------------------------------------------

    clean_text = sanitize_text(text)

    blocks = clean_text.split(
        "\n\n"
    )

    for block in blocks:

        block = block.strip()

        if not block:

            continue

        lines = block.splitlines()

        first_line = (
            lines[0].strip()
        )

        # Detect common headings.
        is_heading = (
            len(first_line) < 120
            and (
                first_line[:2].isdigit()
                or first_line.upper()
                == first_line
                or first_line.endswith(":")
            )
        )

        if is_heading:

            paragraph = (
                document.add_paragraph()
            )

            run = paragraph.add_run(
                first_line
            )

            run.bold = True

            run.font.name = (
                "Times New Roman"
            )

            run.font.size = Pt(12)

            for line in lines[1:]:

                line = line.strip()

                if line:

                    paragraph = (
                        document.add_paragraph(
                            line
                        )
                    )

                    paragraph.style = (
                        document.styles["Normal"]
                    )

        else:

            for line in lines:

                line = line.strip()

                if not line:

                    continue

                paragraph = (
                    document.add_paragraph(
                        line
                    )
                )

                paragraph.style = (
                    document.styles["Normal"]
                )

    # ------------------------------------------------------
    # TERMS TABLE
    # ------------------------------------------------------

    _add_terms_table(
        document,
        terms
    )

    # ------------------------------------------------------
    # FOOTER
    # ------------------------------------------------------

    _add_footer(
        document
    )

    # ------------------------------------------------------
    # SAVE TO MEMORY
    # ------------------------------------------------------

    buffer = BytesIO()

    document.save(
        buffer
    )

    return buffer.getvalue()